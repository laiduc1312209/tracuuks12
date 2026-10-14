# -*- coding: utf-8 -*-
"""Trich du lieu tu TONG-HOP-KQ-KS-LOP12.xlsx ra data.json cho web.

Dung: python tools/extract_web_data.py [dot]   (mac dinh dot = 1)
Doc dotks/<dot>/TONG-HOP-KQ-KS-LOP12.xlsx, ghi web-src/data.json.
"""
import json
import os
import re
import sys
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # thu muc KS
DOT = sys.argv[1] if len(sys.argv) > 1 else "1"                     # dot khao sat
SRC = os.path.join(ROOT, "dotks", DOT, "TONG-HOP-KQ-KS-LOP12.xlsx")
OUT = os.path.join(ROOT, "web-src", "data.json")
print("nguon:", SRC)

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["TongHop"]
hdr = [c.value for c in ws[2]]
subjects = hdr[5:10]                       # Toán, Lý, Hóa, Anh, Văn

students = []
for r in ws.iter_rows(min_row=3, values_only=True):
    if not r[1]:
        continue
    sbd = str(r[1])
    name = re.sub(r"\s+", " ", str(r[2]).strip())
    dob = str(r[3] or "").strip()
    cls = str(r[4] or "").strip()
    scores = []
    for j in range(5):
        v = r[5 + j]
        scores.append(float(v) if isinstance(v, (int, float)) else None)
    students.append({"sbd": sbd, "ten": name, "ns": dob, "lop": cls, "d": scores})

wb.close()

classes = sorted({s["lop"] for s in students if s["lop"]})
payload = {
    "ky": "Khảo sát chất lượng khối 12",
    "mon": subjects,
    "lop": classes,
    "hs": students,
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))

print("hoc sinh:", len(students))
print("lop:", classes)
print("mon:", subjects)
missing = {s: 0 for s in subjects}
for s in students:
    for j, m in enumerate(subjects):
        if s["d"][j] is None:
            missing[m] += 1
print("khong thi theo mon:", missing)
print("size:", os.path.getsize(OUT), "bytes")
print("->", OUT)
